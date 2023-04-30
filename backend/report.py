import psycopg2
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import xlwt
from xlutils.copy import copy
from xlrd import open_workbook
import datetime

conn = psycopg2.connect(dbname='postgres', user='sysdba',
    password='m121144169', host='mk.bmstu-kaluga.ru')
#conn = psycopg2.connect(dbname='mk_db', user='postgres',
#    password='m121144169', host='localhost')
dir = '/usr/share/nginx/html/file_storage/prot/'

cursor = conn.cursor()
cursor.execute("SELECT committee_date, committee_number "+
	" FROM committee "+
	" WHERE committee_id = %(committee_id)s ", {"committee_id": int(sys.argv[1])})
data = cursor.fetchone()
date = data[0]
seria = data[1]
cursor.close()

developments = []
document = Document('./templates/protocol_1.docx')
style = document.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Создаем книку
wb = xlwt.Workbook('utf8')
sheet = wb.add_sheet('Список МУ', cell_overwrite_ok=True)
sheet.col(0).width = 1200
sheet.col(1).width = 3300
sheet.col(2).width = 3000
sheet.col(3).width = 6000
sheet.col(4).width = 24000
sheet.col(5).width = 2400
sheet.col(6).width = 5400

# Стиль шапки
fontAverage = xlwt.Font()
fontAverage.name = 'Times New Roman'
fontAverage.height = 280
fontAverage.bold = True
alignment = xlwt.Alignment()
alignment.horz = xlwt.Alignment.HORZ_CENTER
styleAverage = xlwt.XFStyle()
styleAverage.alignment = alignment
styleAverage.font = fontAverage
borders = xlwt.Borders()
borders.left = 1
borders.right = 1
borders.top = 1
borders.bottom = 1
styleAverage.borders = borders

sheet.write_merge(0, 0, 0, 6, 'Список методических разработок преподавателей для утверждения на заседании Методической комиссии', styleAverage)
sheet.write(1, 0, '№', styleAverage)
sheet.write(1, 1, 'Дата', styleAverage)
sheet.write(1, 2, 'Кафедра', styleAverage)
sheet.write(1, 3, 'Авторы', styleAverage)
sheet.write(1, 4, 'Наименование и вид', styleAverage)
sheet.write(1, 5, 'Стр.', styleAverage)
sheet.write(1, 6, 'Рецензент', styleAverage)

# Стиль текста
fontText = xlwt.Font()
fontText.name = 'Times New Roman'
fontText.height = 240
fontText.bold = False
alignment = xlwt.Alignment()
alignment.horz = xlwt.Alignment.HORZ_CENTER
alignment.vert = xlwt.Alignment.VERT_CENTER
alignment.wrap = xlwt.Alignment.WRAP_AT_RIGHT
styleText = xlwt.XFStyle()
styleText.alignment = alignment
styleText.font = fontText
styleText.borders = borders

dateFormat = xlwt.XFStyle()
dateFormat.num_format_str = 'dd.mm.yyyy'
dateFormat.font = fontText
dateFormat.borders = borders
dateFormat.alignment = alignment

# Стиль текста
fontText = xlwt.Font()
fontText.name = 'Times New Roman'
fontText.height = 240
fontText.bold = False
alignment = xlwt.Alignment()
alignment.horz = xlwt.Alignment.HORZ_LEFT
alignment.vert = xlwt.Alignment.VERT_CENTER
alignment.wrap = xlwt.Alignment.WRAP_AT_RIGHT
nameText = xlwt.XFStyle()
nameText.alignment = alignment
nameText.font = fontText
nameText.borders = borders


cursor2 = conn.cursor()
cursor2.execute("SELECT development_id, development_name, "+
	" development_reviewer, development_pages "+
	" FROM development "+
	" WHERE development_committee_id = %(committee_id)s AND development_state = 3 AND NOT development_review IS NULL "+
	" AND NOT development_cath_resolution IS NULL AND NOT development_deanary_resolution IS NULL"+
	" AND NOT development_review2 IS NULL "+
	" ORDER BY development_id DESC", {"committee_id": sys.argv[1]})
records = cursor2.fetchall()
data = []
for record in records:
	cursor3 = conn.cursor()
	caths = ""
	cath_qountity = 0
	authors_res = ""
	cursor3.execute("SELECT author_id, author_short_name, cath_index "+
		" FROM author INNER JOIN a2dev on a2dev_author_id = author_id " +
		" INNER JOIN cath on cath_id = author_cath_id " +
		" WHERE a2dev_development_id = %(development_id)s", {"development_id": record[0]})
	authors = cursor3.fetchall()
	cursor3.close()
	ind = 0
	for author in authors:
		if ind > 0:
			authors_res = authors_res + ", "
		authors_res = authors_res + str(author[1])
		if caths.find(author[2]) == -1:
			if cath_qountity > 0:
				caths = caths + '; '
			caths = caths + author[2]
			cath_qountity += 1
		ind += 1

	data.append({'id': record[0], 'caths': caths, 'authors': authors_res, 'name': record[1],
        'pages': str(record[3]), 'reviewer': record[2]})

data.sort(key=lambda x: (x['authors'], x['name']))
paragraph_index = 1
for record in data:
    number = str(paragraph_index // 100) + str((paragraph_index - paragraph_index // 100 * 100) // 10) + str(paragraph_index - paragraph_index // 10 * 10) + "/" + seria
    cursor4 = conn.cursor()
    cursor4.execute('UPDATE development SET development_number = %(number)s ' +
		' WHERE development_id = %(id)s', {'number': number, 'id': record['id']})
    cursor4.close()
    conn.commit()

    string = str(paragraph_index) + '. ' + record['authors'] + ' ' + record['name'] + ".- " + record['pages'] + "с. Рег. номер: " + number
    paragraph = document.add_paragraph(string)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph_index += 1
document.save(dir + sys.argv[1] + '.docx')

data.sort(key=lambda x: (x['caths'], x['authors']))
paragraph_index = 1
for record in data:
	sheet.write(paragraph_index + 1, 0, paragraph_index, styleText)
	sheet.write(paragraph_index + 1, 1, date, dateFormat)
	sheet.write(paragraph_index + 1, 2, record['caths'], styleText)
	sheet.write(paragraph_index + 1, 3, record['authors'], styleText)
	sheet.write(paragraph_index + 1, 4, record['name'], nameText)
	sheet.write(paragraph_index + 1, 5, record['pages'], styleText)
	sheet.write(paragraph_index + 1, 6, record['reviewer'], styleText)
	paragraph_index += 1
wb.save(dir + sys.argv[1] + '.xls')
cursor2.close()
conn.close()
